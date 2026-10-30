from __future__ import annotations

import logging
import os
import random
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass
from datetime import date

from build_tracker import JOBS  # single source of truth for posting content
from openpyxl import load_workbook

TRACKER_FILE = "Indeed_Job_Posting_Tracker.xlsx"
TRACKER_SHEET = "Tracker"
FIRST_ROW = 6  # first data row in the Tracker sheet
COL = {"id": 1, "date": 3, "ref": 4, "status": 5, "approver": 6, "proofread": 7, "remarks": 8}

REQUIRED_FIELDS = (
    "title", "type", "location", "experience",
    "qualification", "salary", "skills", "summary", "responsibilities",
)

log = logging.getLogger("indeed")


# --------------------------------------------------------------------------- config
@dataclass(frozen=True)
class Config:
    mode: str  # "mock" or "live"
    client_id: str
    client_secret: str
    employer_id: str

    @classmethod
    def from_env(cls) -> Config:
        return cls(
            mode=os.getenv("INDEED_MODE", "mock").lower(),
            client_id=os.getenv("INDEED_CLIENT_ID", ""),
            client_secret=os.getenv("INDEED_CLIENT_SECRET", ""),
            employer_id=os.getenv("INDEED_EMPLOYER_ID", ""),
        )

    def validate(self) -> None:
        if self.mode not in ("mock", "live"):
            raise ValueError("INDEED_MODE must be 'mock' or 'live'")
        if self.mode == "live":
            missing = [n for n, v in (
                ("INDEED_CLIENT_ID", self.client_id),
                ("INDEED_CLIENT_SECRET", self.client_secret),
                ("INDEED_EMPLOYER_ID", self.employer_id)) if not v]
            if missing:
                raise ValueError(f"Live mode needs: {', '.join(missing)}")


# --------------------------------------------------------------------------- publishers
@dataclass
class PublishResult:
    reference: str   # Indeed job key/URL (live) or SIMULATED-... id (mock)
    status: str      # tracker status value
    simulated: bool
    note: str = ""


class Publisher(ABC):
    @abstractmethod
    def publish(self, job: dict) -> PublishResult: ...


class MockPublisher(Publisher):
    """Pretends to publish. Produces clearly labelled simulated references, never Indeed URLs."""

    def publish(self, job: dict) -> PublishResult:
        ref = f"SIMULATED-{job['id']}-{random.randint(1000, 9999)}"
        # Status stays 'Draft' so simulated runs never count toward the 'Published' target.
        return PublishResult(ref, "Draft", True, "Simulated run - no real posting created.")


class IndeedLivePublisher(Publisher):
    """Skeleton for a real account. Fill in the TODOs from Indeed's partner documentation."""

    MAX_RETRIES = 3

    def __init__(self, cfg: Config):
        self.cfg = cfg
        self._token: str | None = None

    def _authenticate(self) -> str:
        # TODO: exchange client_id/client_secret for an OAuth access token using the
        #       token endpoint from Indeed's partner docs; cache until expiry.
        raise NotImplementedError("Implement OAuth token exchange for your approved Indeed integration.")

    def _build_payload(self, job: dict) -> dict:
        # TODO: map our fields to the schema your integration requires.
        return {
            "employerId": self.cfg.employer_id,
            "title": job["title"],
            "description": f"{job['summary']}\n\nResponsibilities:\n{job['responsibilities']}\n\n"
                           f"Skills: {job['skills']}\nQualification: {job['qualification']}\n"
                           f"Experience: {job['experience']}",
            "location": job["location"],
            "employmentType": job["type"],
            "salary": job["salary"],
        }

    def _send(self, payload: dict) -> str:
        # TODO: POST payload to the job-creation endpoint with the bearer token;
        #       return the Indeed job key / posting URL from the response.
        raise NotImplementedError("Implement the job-creation request for your approved Indeed integration.")

    def publish(self, job: dict) -> PublishResult:
        payload = self._build_payload(job)
        for attempt in range(1, self.MAX_RETRIES + 1):
            try:
                if not self._token:
                    self._token = self._authenticate()
                ref = self._send(payload)
                return PublishResult(ref, "Published", False)
            except NotImplementedError:
                raise
            except Exception as exc:  # network/API errors: retry with backoff
                log.warning("Attempt %d/%d failed for %s: %s", attempt, self.MAX_RETRIES, job["id"], exc)
                if attempt == self.MAX_RETRIES:
                    raise
                time.sleep(2 ** attempt)
        raise RuntimeError("unreachable")


def get_publisher(cfg: Config) -> Publisher:
    return MockPublisher() if cfg.mode == "mock" else IndeedLivePublisher(cfg)


# --------------------------------------------------------------------------- validation
def validate_job(job: dict, live: bool, approver: str | None, proofread: str | None) -> list[str]:
    """Return a list of problems; empty list means OK to publish."""
    problems = [f"missing field: {f}" for f in REQUIRED_FIELDS if not str(job.get(f, "")).strip()]
    if live:
        if "confirm" in job["salary"].lower():
            problems.append("salary still a placeholder - replace with the approved range")
        if not approver:
            problems.append("'Salary Approved By' is empty in the tracker")
        if (proofread or "").upper() != "Y":
            problems.append("posting not marked proofread (Y)")
    return problems


# --------------------------------------------------------------------------- tracker I/O
def update_tracker(results: dict[str, PublishResult], notes: dict[str, str]) -> None:
    wb = load_workbook(TRACKER_FILE)  # formulas preserved (no data_only)
    ws = wb[TRACKER_SHEET]
    for i, job in enumerate(JOBS):
        row = FIRST_ROW + i
        jid = job["id"]
        if jid in results:
            res = results[jid]
            ws.cell(row=row, column=COL["date"], value=date.today()).number_format = "dd-mmm-yyyy"
            ws.cell(row=row, column=COL["ref"], value=res.reference)
            ws.cell(row=row, column=COL["status"], value=res.status)
            ws.cell(row=row, column=COL["remarks"], value=res.note or "Published via API.")
        elif jid in notes:
            ws.cell(row=row, column=COL["remarks"], value=notes[jid])
    wb.save(TRACKER_FILE)


def read_tracker_state() -> dict[str, dict]:
    wb = load_workbook(TRACKER_FILE)
    ws = wb[TRACKER_SHEET]
    state = {}
    for i, job in enumerate(JOBS):
        row = FIRST_ROW + i
        state[job["id"]] = {
            "status": ws.cell(row=row, column=COL["status"]).value,
            "approver": ws.cell(row=row, column=COL["approver"]).value,
            "proofread": ws.cell(row=row, column=COL["proofread"]).value,
        }
    return state


# --------------------------------------------------------------------------- main
def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    cfg = Config.from_env()
    cfg.validate()
    live = cfg.mode == "live"
    publisher = get_publisher(cfg)
    log.info("Mode: %s", cfg.mode.upper())

    state = read_tracker_state()
    results: dict[str, PublishResult] = {}
    notes: dict[str, str] = {}

    for job in JOBS:
        jid = job["id"]
        if state[jid]["status"] == "Published":  # idempotent: never double-post
            log.info("%s already published - skipped", jid)
            continue
        problems = validate_job(job, live, state[jid]["approver"], state[jid]["proofread"])
        if problems:
            log.error("%s blocked: %s", jid, "; ".join(problems))
            notes[jid] = "Blocked: " + "; ".join(problems)
            continue
        try:
            results[jid] = publisher.publish(job)
            log.info("%s -> %s (%s)", jid, results[jid].reference, results[jid].status)
        except NotImplementedError as exc:
            log.error("%s: %s", jid, exc)
            notes[jid] = f"Live publisher not implemented yet: {exc}"
        # except Exception as exc:
        #     log.error("%s failed: %s", jid, exc)
        #     notes[jid] = f"Publish failed: {exc}"

    update_tracker(results, notes)
    log.info("Tracker updated: %d published/simulated, %d not processed.",
             len(results), len(JOBS) - len(results))


if __name__ == "__main__":
    main()
