from __future__ import annotations

import csv
import logging
import os
from abc import ABC, abstractmethod
from datetime import date, timedelta

from build_tracker import JOBS
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

log = logging.getLogger("applicants")
OUTPUT_FILE = "Indeed_Applicant_Summary.xlsx"

# ----------------------------------------------------------------------------- rules
JOB_CRITERIA = {
    "JP-001": {
        "skills": ["python", "django", "fastapi", "rest api", "postgresql", "docker", "git", "aws"],
        "min_exp": 4, "education": ["b.tech", "b.e", "mca", "m.tech", "computer science"],
        "locations": ["gorakhpur", "lucknow", "remote"], "max_notice_days": 60, "salary_max_lpa": 20,
    },
    "JP-002": {
        "skills": ["react", "javascript", "typescript", "html", "css", "redux", "rest api", "git"],
        "min_exp": 2, "education": ["b.tech", "b.e", "bca", "mca", "computer science"],
        "locations": ["gorakhpur"], "max_notice_days": 45, "salary_max_lpa": 10,
    },
    "JP-003": {
        "skills": ["selenium", "playwright", "python", "java", "postman", "jenkins", "sql", "test planning"],
        "min_exp": 2, "education": ["b.tech", "b.e", "mca", "computer science"],
        "locations": ["gorakhpur", "lucknow"], "max_notice_days": 45, "salary_max_lpa": 9,
    },
}
REQUIRED_PROFILE_FIELDS = ("skills", "experience_years", "education", "location")
CATEGORIES = ["Relevant/Suitable", "Partially Relevant", "Not Relevant", "Incomplete Profile"]
EXP_LEVELS = ["Experienced", "Fresher"]
FRESHER_BELOW_YEARS = 1          # assumption: < 1 year of relevant experience = Fresher
RELEVANT_SCORE, PARTIAL_SCORE = 70, 40   # assumption: score thresholds (documented in workbook)
WEIGHTS = {"skills": 50, "experience": 25, "education": 10, "location": 10, "availability_salary": 5}


# ----------------------------------------------------------------------------- sources
class ApplicantSource(ABC):
    @abstractmethod
    def fetch(self) -> list[dict]: ...


class SampleSource(ApplicantSource):
    """Illustrative SAMPLE applicants so the pipeline can be demonstrated without a live account."""

    def fetch(self) -> list[dict]:
        d = date.today()
        rows = [
            # job, name, email, location, exp, skills, education, expected LPA, notice days, days ago
            ("JP-001", "Sample Applicant 01", "s01@example.com", "Gorakhpur", 6, "Python;Django;PostgreSQL;Docker;AWS;Git;REST API", "B.Tech Computer Science", 18, 30, 1),
            ("JP-001", "Sample Applicant 02", "s02@example.com", "Lucknow", 4, "Python;FastAPI;Git;REST API", "MCA", 19, 60, 2),
            ("JP-001", "Sample Applicant 03", "s03@example.com", "Delhi", 2, "Java;Spring", "B.Tech", 14, 30, 2),
            ("JP-001", "Sample Applicant 04", "s04@example.com", "Gorakhpur", 0, "Python;Git", "B.Tech Computer Science", 5, 0, 3),
            ("JP-001", "Sample Applicant 05", "s05@example.com", "", None, "", "", None, None, 3),
            ("JP-002", "Sample Applicant 06", "s06@example.com", "Gorakhpur", 3, "React;JavaScript;TypeScript;Redux;HTML;CSS;Git", "BCA", 9, 30, 1),
            ("JP-002", "Sample Applicant 07", "s07@example.com", "Gorakhpur", 0, "HTML;CSS;JavaScript", "B.Tech Computer Science", 4, 0, 2),
            ("JP-002", "Sample Applicant 08", "s08@example.com", "Mumbai", 5, "Angular;JavaScript", "B.E", 16, 90, 4),
            ("JP-002", "Sample Applicant 09", "s09@example.com", "Gorakhpur", 2, "React;JavaScript;CSS", "B.Com", 8, 15, 4),
            ("JP-002", "Sample Applicant 10", "s10@example.com", "Gorakhpur", 4, "Accounting;Tally", "B.Com", 4, 15, 5),
            ("JP-003", "Sample Applicant 11", "s11@example.com", "Gorakhpur", 3, "Selenium;Python;Postman;Jenkins;SQL;Test Planning", "B.Tech Computer Science", 8, 30, 1),
            ("JP-003", "Sample Applicant 12", "s12@example.com", "Lucknow", 2, "Selenium;Java", "MCA", 7, 45, 2),
            ("JP-003", "Sample Applicant 13", "s13@example.com", "Gorakhpur", 0, "Manual Testing;SQL", "B.Tech", 3.5, 0, 3),
            ("JP-003", "Sample Applicant 14", "s14@example.com", "Pune", 4, "Playwright;Python;SQL", "B.E", 12, 30, 5),
            ("JP-003", "Sample Applicant 15", "s15@example.com", "Gorakhpur", 3, "Selenium;Python;SQL", "B.Tech", 8, 30, 5),
            ("JP-003", "Sample Applicant 15", "s15@example.com", "Gorakhpur", 3, "Selenium;Python;SQL", "B.Tech", 8, 30, 5),  # duplicate
        ]
        keys = ("job_id", "name", "email", "location", "experience_years", "skills",
                "education", "expected_salary_lpa", "notice_days", "days_ago")
        out = []
        for i, r in enumerate(rows, 1):
            rec = dict(zip(keys, r))
            rec["applied_date"] = d - timedelta(days=rec.pop("days_ago"))
            rec["applicant_id"] = f"AP-{i:03d}"
            out.append(rec)
        return out


class CsvSource(ApplicantSource):
    """Reads a CSV with columns: job_id,name,email,location,experience_years,skills,education,
    expected_salary_lpa,notice_days,applied_date(YYYY-MM-DD). Map Indeed's export headers to these."""

    def __init__(self, path: str):
        self.path = path

    @staticmethod
    def _num(v):
        try:
            return float(v) if str(v).strip() else None
        except ValueError:
            return None

    def fetch(self) -> list[dict]:
        out = []
        with open(self.path, newline="", encoding="utf-8-sig") as f:
            for i, r in enumerate(csv.DictReader(f), 1):
                try:
                    applied = date.fromisoformat(r.get("applied_date", "").strip())
                except ValueError:
                    applied = None
                out.append({
                    "applicant_id": f"AP-{i:03d}", "job_id": r.get("job_id", "").strip(),
                    "name": r.get("name", "").strip(), "email": r.get("email", "").strip().lower(),
                    "phone": r.get("phone", "").strip(),
                    "location": r.get("location", "").strip(), "skills": r.get("skills", "").strip(),
                    "education": r.get("education", "").strip(),
                    "experience_years": self._num(r.get("experience_years")),
                    "expected_salary_lpa": self._num(r.get("expected_salary_lpa")),
                    "notice_days": self._num(r.get("notice_days")), "applied_date": applied,
                })
        return out


def get_source() -> ApplicantSource:
    mode = os.getenv("APPLICANT_SOURCE", "sample").lower()
    if mode == "csv":
        path = os.getenv("APPLICANT_CSV")
        if not path:
            raise ValueError("APPLICANT_SOURCE=csv needs APPLICANT_CSV=<path>")
        return CsvSource(path)
    if mode != "sample":
        raise ValueError("APPLICANT_SOURCE must be 'sample' or 'csv'")
    return SampleSource()


# ----------------------------------------------------------------------------- classification
def _skill_list(text: str) -> set[str]:
    return {s.strip().lower() for s in str(text or "").replace(",", ";").split(";") if s.strip()}


def classify(app: dict) -> dict:
    """Return category, experience level, score and job-related reasons for one applicant."""
    crit = JOB_CRITERIA.get(app["job_id"])
    if crit is None:
        return {"category": "Incomplete Profile", "level": "", "score": 0,
                "reasons": f"Unknown job id '{app['job_id']}'"}

    exp = app.get("experience_years")
    level = "" if exp is None else ("Fresher" if exp < FRESHER_BELOW_YEARS else "Experienced")

    missing = [f for f in REQUIRED_PROFILE_FIELDS if app.get(f) in (None, "")]
    if missing:
        return {"category": "Incomplete Profile", "level": level, "score": 0,
                "reasons": "Missing: " + ", ".join(missing)}

    reasons, score = [], 0.0
    # skills
    have = _skill_list(app["skills"])
    matched = [s for s in crit["skills"] if any(s in h or h in s for h in have)]
    ratio = len(matched) / len(crit["skills"])
    score += WEIGHTS["skills"] * ratio
    reasons.append(f"Skills {len(matched)}/{len(crit['skills'])} matched ({', '.join(matched) or 'none'})")
    # experience
    need = crit["min_exp"]
    exp_ratio = min(exp / need, 1.0)
    score += WEIGHTS["experience"] * exp_ratio
    reasons.append(f"Experience {exp:g} yrs vs {need} required" + ("" if exp >= need else " (below requirement)"))
    # education
    edu_ok = any(e in app["education"].lower() for e in crit["education"])
    score += WEIGHTS["education"] * edu_ok
    reasons.append("Education fits" if edu_ok else "Education not in preferred list")
    # location
    loc_ok = any(l in app["location"].lower() for l in crit["locations"])
    score += WEIGHTS["location"] * loc_ok
    reasons.append(f"Location {app['location']} " + ("acceptable" if loc_ok else "outside job locations"))
    # availability + salary (only when provided)
    av_pts = WEIGHTS["availability_salary"]
    notes = []
    notice = app.get("notice_days")
    if notice is not None:
        ok = notice <= crit["max_notice_days"]
        notes.append(f"Notice {notice:g}d " + ("OK" if ok else "too long"))
        av_pts -= 0 if ok else av_pts / 2
    sal = app.get("expected_salary_lpa")
    if sal is not None:
        ok = sal <= crit["salary_max_lpa"]
        notes.append(f"Expected {sal:g} LPA " + ("within" if ok else "above") + f" approved max {crit['salary_max_lpa']}")
        av_pts -= 0 if ok else av_pts / 2
    score += max(av_pts, 0)
    reasons.extend(notes)

    score = round(score)
    category = ("Relevant/Suitable" if score >= RELEVANT_SCORE
                else "Partially Relevant" if score >= PARTIAL_SCORE else "Not Relevant")
    if not matched:  # experience in an unrelated field must not make someone relevant
        category = "Not Relevant"
        reasons.append("No required skills matched - experience is not in the job's field")
    return {"category": category, "level": level, "score": score, "reasons": "; ".join(reasons)}


def dedupe(apps: list[dict]) -> tuple[list[dict], int]:
    seen, out, dups = set(), [], 0
    for a in apps:
        key = (a["job_id"], (a.get("email") or a.get("name") or "").lower())
        if key in seen:
            dups += 1
            continue
        seen.add(key)
        out.append(a)
    return out, dups


# ----------------------------------------------------------------------------- workbook
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

APP_COLS = [("Applicant ID", 12), ("Applicant (name)", 22), ("Job ID", 9), ("Job Title", 34),
            ("Applied Date", 13), ("Location", 14), ("Experience (yrs)", 12), ("Education", 24),
            ("Skills", 36), ("Expected Salary (LPA)", 12), ("Notice (days)", 10),
            ("Match Score (0-100)", 12), ("Category", 20), ("Experience Level", 16),
            ("Reasons (job-related)", 70), ("Reviewer Override / Notes", 28)]


def _style_header(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BORDER


def write_workbook(rows: list[dict], dup_count: int) -> None:
    titles = {j["id"]: j["title"] for j in JOBS}
    wb = Workbook()
    ws = wb.active
    ws.title = "Applicants"
    for c, (h, w) in enumerate(APP_COLS, 1):
        ws.cell(row=1, column=c, value=h)
        ws.column_dimensions[get_column_letter(c)].width = w
    _style_header(ws, 1, len(APP_COLS))
    for r, a in enumerate(sorted(rows, key=lambda x: x["applicant_id"]), 2):
        vals = [a["applicant_id"], a["name"], a["job_id"], titles.get(a["job_id"], "Unknown"),
                a["applied_date"], a["location"], a["experience_years"], a["education"], a["skills"],
                a["expected_salary_lpa"], a["notice_days"], a["score"], a["category"], a["level"],
                a["reasons"], None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.cell(row=r, column=5).number_format = "dd-mmm-yyyy"
        ws.row_dimensions[r].height = 45
    last = 1 + len(rows)
    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=False)
    dv_lvl = DataValidation(type="list", formula1='"' + ",".join(EXP_LEVELS) + '"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_lvl)
    dv_cat.add(f"M2:M{last}")
    dv_lvl.add(f"N2:N{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(APP_COLS))}{last}"

    # ---- Summary (formulas so reviewer overrides flow through)
    s = wb.create_sheet("Summary", 0)
    s.column_dimensions["A"].width = 38
    for col in "BCDEFGH":
        s.column_dimensions[col].width = 16
    s["A1"] = "Indeed Applications - Summary"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    s["A2"] = f"Generated {date.today():%d-%b-%Y}. Data source: sample/export (see Notes). Duplicates removed: {dup_count}."
    s["A2"].font = Font(name=FONT, italic=True, size=9)
    s["A4"], s["B4"] = "Total applications", f"=COUNTA(Applicants!A2:A{last})"
    s["A4"].font = Font(name=FONT, bold=True)
    s["B4"].font = Font(name=FONT, bold=True)

    # Category counts
    s["A6"], s["B6"], s["C6"] = "Category", "Count", "% of total"
    _style_header(s, 6, 3)
    r = 7
    for cat in CATEGORIES + EXP_LEVELS:
        col = "M" if cat in CATEGORIES else "N"
        s.cell(row=r, column=1, value=cat)
        s.cell(row=r, column=2, value=f'=COUNTIF(Applicants!{col}2:{col}{last},A{r})')
        s.cell(row=r, column=3, value=f"=IF($B$4=0,0,B{r}/$B$4)").number_format = "0.0%"
        for c in range(1, 4):
            s.cell(row=r, column=c).font = Font(name=FONT)
            s.cell(row=r, column=c).border = BORDER
        r += 1
    s.cell(row=r, column=1, value="Check: suitability categories sum to total").font = Font(name=FONT, italic=True, size=9)
    s.cell(row=r, column=2, value='=IF(SUM(B7:B10)=B4,"OK","MISMATCH")').font = Font(name=FONT, italic=True, size=9)

    # Per-job breakdown
    r += 2
    heads = ["Job"] + CATEGORIES + ["Total"]
    for c, h in enumerate(heads, 1):
        s.cell(row=r, column=c, value=h)
    _style_header(s, r, len(heads))
    for job in JOBS:
        r += 1
        s.cell(row=r, column=1, value=f"{job['id']} - {job['title']}")
        for c, cat in enumerate(CATEGORIES, 2):
            s.cell(row=r, column=c, value=f'=COUNTIFS(Applicants!$C$2:$C${last},"{job["id"]}",Applicants!$M$2:$M${last},"{cat}")')
        s.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})")
        for c in range(1, 7):
            s.cell(row=r, column=c).font = Font(name=FONT)
            s.cell(row=r, column=c).border = BORDER
            s.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    # Notes / assumptions
    r += 2
    notes = [
        "Notes & assumptions",
        f"- Score = skills {WEIGHTS['skills']} + experience {WEIGHTS['experience']} + education {WEIGHTS['education']} + location {WEIGHTS['location']} + notice/salary {WEIGHTS['availability_salary']}.",
        f"- Relevant/Suitable: score >= {RELEVANT_SCORE}; Partially Relevant: >= {PARTIAL_SCORE}; otherwise Not Relevant.",
        "- An applicant matching none of the required skills is always Not Relevant, whatever the score.",
        "- Incomplete Profile: skills, experience, education or location missing (cannot be scored fairly).",
        f"- Fresher: under {FRESHER_BELOW_YEARS} year of experience; Experienced: {FRESHER_BELOW_YEARS} year or more.",
        "- Only job-related criteria are used. A recruiter should review borderline cases; edit Category/Experience Level on the Applicants sheet and this summary updates.",
        "- Applicants shown are SAMPLE records unless APPLICANT_SOURCE=csv was used.",
    ]
    for i, t in enumerate(notes):
        c = s.cell(row=r + i, column=1, value=t)
        c.font = Font(name=FONT, bold=(i == 0), size=10 if i else 11)
    wb.save(OUTPUT_FILE)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    raw = get_source().fetch()
    apps, dups = dedupe(raw)
    log.info("Fetched %d applications, %d duplicates removed", len(raw), dups)
    for a in apps:
        a.update(classify(a))
    write_workbook(apps, dups)
    counts = {c: sum(1 for a in apps if a["category"] == c) for c in CATEGORIES}
    log.info("Total %d | %s", len(apps), " | ".join(f"{k}: {v}" for k, v in counts.items()))
    log.info("Saved %s", OUTPUT_FILE)


if __name__ == "__main__":
    main()
